from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from docuflow.ai.chunking import chunk_pages, retrieve
from docuflow.ai.providers import MockAIProvider
from docuflow.conversion.engines import ConversionService, write_simple_docx
from docuflow.pdf.engine import PdfEngine
from tests.pdf.conftest import write_text_pdf


def test_mock_summary_uses_real_pages(tmp_path: Path) -> None:
    pdf = write_text_pdf(tmp_path / "ai.pdf", ["Quarterly revenue grew in Jakarta.", "Hiring plan for Bandung office."])
    extracted = PdfEngine().extract_text(pdf)
    chunks = chunk_pages(extracted.metadata["pages"])
    result = MockAIProvider().generate(instruction="Summarize", chunks=chunks)
    cited_pages = {item["page"] for item in result.citations}
    assert cited_pages.issubset({1, 2})
    assert "Jakarta" in result.text or "Bandung" in result.text or "Page" in result.text


def test_retrieval_ranks_relevant_chunk() -> None:
    pages = [
        {"page": 1, "text": "Apples are harvested in autumn."},
        {"page": 2, "text": "The contract value is 12 billion rupiah."},
        {"page": 3, "text": "Unrelated weather notes."},
    ]
    chunks = chunk_pages(pages)
    found = retrieve(chunks, "contract value rupiah")
    assert found[0].page == 2


def test_pdf_to_xlsx_exports_tables(tmp_path: Path) -> None:
    import fitz

    path = tmp_path / "table.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Name  Amount")
    page.insert_text((72, 92), "Ada   10")
    page.insert_text((72, 112), "Grace 20")
    doc.save(path.as_posix())
    doc.close()
    dest = tmp_path / "out.xlsx"
    result = ConversionService().pdf_to_xlsx(path, dest)
    assert dest.is_file()
    assert result.output_format == "xlsx"
    workbook = load_workbook(dest)
    assert workbook.sheetnames


def test_pdf_to_docx(tmp_path: Path) -> None:
    pdf = write_text_pdf(tmp_path / "letter.pdf", ["Dear team, please find the report attached."])
    dest = tmp_path / "letter.docx"
    result = ConversionService().pdf_to_docx(pdf, dest)
    assert dest.is_file()
    assert result.output_format == "docx"
    assert result.limitations


def test_docx_fixture_roundtrip_kind(tmp_path: Path) -> None:
    from docuflow.pdf.sniff import sniff_file

    path = write_simple_docx(tmp_path / "note.docx", ["Hello from DocuFlow"])
    assert sniff_file(path).kind == "docx"


def test_docx_to_pdf_works_without_libreoffice(tmp_path: Path, monkeypatch) -> None:
    import fitz
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    from PIL import Image

    from docuflow.conversion.docx_pdf import convert_docx_with_builtin
    from docuflow.pdf.sniff import sniff_file

    monkeypatch.setattr("docuflow.conversion.engines.find_libreoffice", lambda: None)
    monkeypatch.setattr("docuflow.conversion.docx_pdf.find_libreoffice", lambda: None)

    image_path = tmp_path / "dot.png"
    Image.new("RGB", (32, 32), color=(20, 80, 180)).save(image_path)

    source = tmp_path / "brief.docx"
    document = Document()
    document.add_heading("Quarterly brief", level=1)
    document.add_paragraph("Hello from DocuFlow Word to PDF.")
    centered = document.add_paragraph("Centered note")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Ada"
    table.cell(1, 1).text = "10"
    document.add_picture(str(image_path), width=Inches(1))
    document.save(str(source))

    dest = tmp_path / "brief.pdf"
    result = ConversionService().docx_to_pdf(source, dest)
    assert dest.is_file()
    assert dest.read_bytes()[:5] == b"%PDF-"
    assert sniff_file(dest).kind == "pdf"
    assert result.engine == "docx-builtin"
    assert result.output_format == "pdf"

    pdf = fitz.open(dest.as_posix())
    text = "".join(page.get_text() for page in pdf)
    pdf.close()
    assert "Quarterly brief" in text
    assert "Hello from DocuFlow" in text
    assert "Ada" in text

    convert_docx_with_builtin(source, tmp_path / "again.pdf")
    assert (tmp_path / "again.pdf").is_file()


def test_docx_to_pdf_falls_back_when_libreoffice_fails(tmp_path: Path, monkeypatch) -> None:
    from docuflow.exceptions import EngineUnavailableError

    source = write_simple_docx(tmp_path / "note.docx", ["Fallback after LibreOffice failure."])
    dest = tmp_path / "note.pdf"

    def boom(_self: ConversionService, _source: Path, _output: Path, _soffice: str) -> None:
        raise EngineUnavailableError("LibreOffice could not convert the document.")

    monkeypatch.setattr("docuflow.conversion.engines.find_libreoffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(ConversionService, "_docx_to_pdf_libreoffice", boom)
    result = ConversionService().docx_to_pdf(source, dest)
    assert dest.is_file()
    assert result.engine == "docx-builtin"
