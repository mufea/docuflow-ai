from __future__ import annotations

import logging
import os
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document as DocxDocument
from openpyxl import Workbook
from pdf2docx import Converter

from docuflow.conversion.docx_pdf import (
    BUILTIN_DOCX_TO_PDF_LIMITS,
    convert_docx_with_builtin,
    docx_to_pdf_available,
    find_libreoffice,
)
from docuflow.exceptions import EngineUnavailableError, InvalidFileError, UnsupportedOperationError
from docuflow.pdf.sniff import sniff_file

logger = logging.getLogger("docuflow.conversion")


@dataclass
class ConversionResult:
    outputs: list[Path]
    output_format: str
    engine: str
    limitations: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


PDF_TO_DOCX_LIMITS = [
    "Layout reconstruction is approximate. Multi-column pages, footnotes, and complex tables may shift.",
    "Scanned pages without a text layer will produce an empty or sparse DOCX unless OCR is run first.",
    "Fonts, headers, and precise positioning are not guaranteed.",
]

PDF_TO_XLSX_LIMITS = [
    "Only detected tables are exported. Free-form text is ignored.",
    "Merged cells, nested tables, and ruled-line tables without extractable structure may be incomplete.",
]

DOCX_TO_PDF_LIMITS = [
    "LibreOffice is used when installed. Otherwise a built-in converter writes an approximate PDF.",
    "Pagination, fonts, headers, and floating objects may differ from Word.",
    "Macros are never executed. Embedded scripts are not processed.",
]


class ConversionService:
    def pdf_to_docx(self, source: Path, output: Path) -> ConversionResult:
        kind = sniff_file(source)
        if kind.kind != "pdf":
            raise UnsupportedOperationError("PDF to DOCX requires a PDF input.")
        output.parent.mkdir(parents=True, exist_ok=True)
        converter = Converter(source.as_posix())
        try:
            converter.convert(output.as_posix())
        finally:
            converter.close()
        if not output.is_file() or output.stat().st_size == 0:
            raise InvalidFileError("DOCX conversion produced an empty file.")
        return ConversionResult(
            outputs=[output],
            output_format="docx",
            engine="pdf2docx",
            limitations=PDF_TO_DOCX_LIMITS,
        )

    def pdf_to_xlsx(self, source: Path, output: Path) -> ConversionResult:
        kind = sniff_file(source)
        if kind.kind != "pdf":
            raise UnsupportedOperationError("PDF to XLSX requires a PDF input.")
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        tables_found = 0
        with pdfplumber.open(source) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                if not tables:
                    continue
                for table_index, table in enumerate(tables, start=1):
                    sheet = workbook.create_sheet(title=f"p{index}_t{table_index}"[:31])
                    for row in table:
                        sheet.append([cell if cell is not None else "" for cell in row])
                    tables_found += 1
        if tables_found == 0:
            sheet = workbook.create_sheet(title="no_tables")
            sheet.append(["No tables were detected in this PDF."])
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
        return ConversionResult(
            outputs=[output],
            output_format="xlsx",
            engine="pdfplumber+openpyxl",
            limitations=PDF_TO_XLSX_LIMITS,
            metadata={"tables_found": tables_found},
        )

    def docx_to_pdf(self, source: Path, output: Path) -> ConversionResult:
        kind = sniff_file(source)
        if kind.kind != "docx":
            raise UnsupportedOperationError("DOCX to PDF requires a DOCX input.")
        soffice = find_libreoffice()
        if soffice:
            try:
                self._docx_to_pdf_libreoffice(source, output, soffice)
                return ConversionResult(
                    outputs=[output],
                    output_format="pdf",
                    engine="libreoffice-headless",
                    limitations=DOCX_TO_PDF_LIMITS,
                )
            except EngineUnavailableError:
                logger.warning("libreoffice_failed_using_builtin")
        convert_docx_with_builtin(source, output)
        return ConversionResult(
            outputs=[output],
            output_format="pdf",
            engine="docx-builtin",
            limitations=BUILTIN_DOCX_TO_PDF_LIMITS,
        )

    def _docx_to_pdf_libreoffice(self, source: Path, output: Path, soffice: str) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        env = {
            "PATH": os.environ.get("PATH", ""),
            "HOME": output.parent.as_posix(),
            "LANG": "C",
        }
        try:
            completed = subprocess.run(  # noqa: S603
                [
                    soffice,
                    "--headless",
                    "--norestore",
                    "--nolockcheck",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--convert-to",
                    "pdf:writer_pdf_Export",
                    "--outdir",
                    str(output.parent),
                    str(source),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=120,
                env=env,
                cwd=output.parent,
            )
        except subprocess.TimeoutExpired as exc:
            raise EngineUnavailableError("LibreOffice conversion timed out.") from exc
        produced = source.with_suffix(".pdf")
        candidate = output.parent / produced.name
        if completed.returncode != 0 or not candidate.is_file():
            logger.warning("libreoffice_failed code=%s", completed.returncode)
            raise EngineUnavailableError("LibreOffice could not convert the document.")
        if candidate != output:
            candidate.replace(output)

    def docx_preview_unsupported_without_engine(self) -> dict[str, Any]:
        return {
            "operation": "docx_to_pdf",
            "available": docx_to_pdf_available(),
            "engine": "libreoffice-headless" if find_libreoffice() else "docx-builtin",
            "limitations": DOCX_TO_PDF_LIMITS,
        }


def write_simple_docx(path: Path, paragraphs: list[str]) -> Path:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(str(path))
    return path
